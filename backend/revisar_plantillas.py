
import glob
import os
import openpyxl

carpeta = "app/templates/informes_excel"
archivos = glob.glob(os.path.join(carpeta, "*.xlsx"))

print("ARCHIVOS ENCONTRADOS:", len(archivos))

for archivo in archivos:
    print("\n" + "=" * 60)
    print("ARCHIVO:", os.path.basename(archivo))

    libro = openpyxl.load_workbook(archivo, read_only=False, data_only=True)

    for hoja in libro.worksheets:
        print("\nHOJA:", hoja.title)
        print("FILAS:", hoja.max_row, "| COLUMNAS:", hoja.max_column)

        for numero_fila in (1, 2):
            encabezados = [
                hoja.cell(numero_fila, columna).value
                for columna in range(1, min(hoja.max_column, 30) + 1)
            ]
            print(f"FILA {numero_fila}:", encabezados)

    libro.close()
