import os
import tempfile
import unicodedata
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, joinedload

from app.core.errors import AppError
from app.models.audit import Auditoria
from app.models.commercial import Compra
from app.models.organization import Organizacion
from app.models.user import Usuario
from app.services import panel as panel_service
from app.services.panel import TIPO_POR_ACCION

TABLAS_VALIDAS = {"usuarios", "empresas", "compras", "auditoria"}

TEXTO_ACCION = {
    "LOGIN": "inició sesión",
    "LOGIN_GOOGLE": "inició sesión con Google",
    "CREAR_USUARIO": "creó un usuario",
    "ACTUALIZAR_USUARIO": "actualizó un usuario",
    "DESACTIVAR_USUARIO": "desactivó un usuario",
    "CAMBIAR_ROL_USUARIO": "cambió el rol de un usuario",
    "EDITAR_PERFIL": "editó su perfil",
    "CAMBIAR_PASSWORD_PERFIL": "cambió su contraseña",
    "CREAR_ENLACE_PACIENTES": "creó un enlace de pacientes",
    "ELIMINAR_ENLACE_PACIENTES": "eliminó un enlace de pacientes",
    "CREAR_EVALUACION": "creó una evaluación",
    "INICIAR_EVALUACION": "inició una evaluación",
    "FINALIZAR_EVALUACION": "finalizó una evaluación",
    "INFORME_INDIVIDUAL_GENERADO": "consultó un resultado individual",
    "INFORME_AGRUPADO_GENERADO": "descargó un informe agregado",
    "INFORME_DESCARGADO": "descargó un informe",
    "EDITAR_EMPRESA": "editó una empresa",
    "EDITAR_RESPONSABLE_EMPRESA": "editó el responsable de una empresa",
    "ELIMINAR_EMPRESA": "eliminó una empresa",
    "CAMBIAR_EMPRESA_ACTIVA": "cambió su empresa activa",
    "INVITAR_USUARIO": "invitó a un usuario",
    "GUARDAR_RESOLUCION": "guardó su resolución",
    "VERIFICAR_RESOLUCION": "verificó una resolución",
    "QUITAR_VERIFICACION_RESOLUCION": "quitó la verificación de una resolución",
    "CREAR_PLAN": "creó un plan",
    "EDITAR_PLAN": "editó un plan",
    "ELIMINAR_PLAN": "eliminó un plan",
    "EXPORTAR_PANEL": "exportó una tabla del panel",
}


def _normalizar(txt: str | None) -> str:
    if not txt:
        return ""
    nfkd = unicodedata.normalize("NFKD", str(txt).lower())
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def _formato_fecha(dt) -> str:
    if not dt:
        return "—"
    if isinstance(dt, str):
        try:
            dt = datetime.fromisoformat(dt)
        except Exception:
            return dt
    return dt.strftime("%d/%m/%Y %H:%M")


def exportar_tabla_excel(
    db: Session,
    tabla: str,
    filtro: str | None = None,
    buscar: str | None = None,
) -> str:
    if tabla not in TABLAS_VALIDAS:
        raise AppError(400, "Tabla no válida.")

    wb = Workbook()
    ws = wb.active
    ws.title = tabla.capitalize()

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    b_norm = _normalizar(buscar.strip()) if buscar else ""

    if tabla == "usuarios":
        cabeceras = [
            "Nombre",
            "Correo",
            "Rol",
            "Empresa",
            "Registro",
            "Estado",
            "Último acceso",
            "Número de resolución",
            "Resolución verificada",
        ]
        ws.append(cabeceras)

        panel_datos = panel_service.obtener_panel(db)
        usuarios = panel_datos["usuarios"]

        if b_norm:
            usuarios = [
                u
                for u in usuarios
                if b_norm in _normalizar(u.get("nombre"))
                or b_norm in _normalizar(u.get("correo"))
                or b_norm in _normalizar(u.get("empresa"))
            ]

        for u in usuarios:
            ultimo_acc = _formato_fecha(u.get("ultimoAcceso"))
            res_verif = "Sí" if u.get("resolucionVerificada") else "No"
            num_res = u.get("numeroResolucion") or "—"
            ws.append([
                u.get("nombre", ""),
                u.get("correo", ""),
                u.get("rol", ""),
                u.get("empresa", "—"),
                u.get("registro", ""),
                u.get("estado", ""),
                ultimo_acc,
                num_res,
                res_verif,
            ])

    elif tabla == "empresas":
        cabeceras = [
            "Empresa",
            "NIT",
            "Responsable SST",
            "Encuestas completadas",
            "Evaluaciones usadas",
            "Evaluaciones compradas",
        ]
        ws.append(cabeceras)

        panel_datos = panel_service.obtener_panel(db)
        empresas = panel_datos["empresas"]

        if b_norm:
            empresas = [
                e
                for e in empresas
                if b_norm in _normalizar(e.get("nombre"))
                or b_norm in _normalizar(e.get("nit"))
                or b_norm in _normalizar(e.get("responsable"))
            ]

        for e in empresas:
            ws.append([
                e.get("nombre", ""),
                e.get("nit", ""),
                e.get("responsable", "—"),
                e.get("encuestas", 0),
                e.get("usadas", 0),
                e.get("compradas", 0),
            ])

    elif tabla == "compras":
        cabeceras = [
            "Fecha",
            "Empresa",
            "Cantidad",
            "Medio de pago",
            "Valor",
            "Estado",
        ]
        ws.append(cabeceras)

        orgs = {o.id: o for o in db.query(Organizacion).all()}
        compras_db = (
            db.query(Compra)
            .options(joinedload(Compra.pagos))
            .order_by(Compra.creado_en.desc())
            .all()
        )

        for c in compras_db:
            pago = max(c.pagos, key=lambda p: p.procesado_en, default=None) if c.pagos else None
            medio = "—"
            if pago:
                medio = pago.metodo + (f" •••• {pago.ultimos_digitos}" if pago.ultimos_digitos else "")
            empresa = orgs[c.organizacion_id].nombre if c.organizacion_id in orgs else "—"
            fecha = _formato_fecha(c.creado_en)
            ws.append([
                fecha,
                empresa,
                c.cantidad,
                medio,
                c.total,
                c.estado,
            ])
            ws.cell(row=ws.max_row, column=5).number_format = '"$"#,##0'

    elif tabla == "auditoria":
        cabeceras = [
            "Fecha",
            "Usuario",
            "Acción",
            "Tipo",
            "Detalle",
        ]
        ws.append(cabeceras)

        usuarios_db = db.query(Usuario).filter(Usuario.es_invitado.is_(False)).all()
        nombres = {u.id: f"{u.nombre} {u.apellido}".strip() for u in usuarios_db}
        registros_aud = db.query(Auditoria).order_by(Auditoria.registrado_en.desc()).all()

        f_tipo = filtro.strip() if filtro else ""

        for a in registros_aud:
            tipo = TIPO_POR_ACCION.get(a.accion, "Usuarios")
            if f_tipo and f_tipo != "Todo" and tipo != f_tipo:
                continue

            u_nombre = nombres.get(a.usuario_id, "Sistema")
            accion_legible = TEXTO_ACCION.get(a.accion, a.accion)
            if b_norm:
                if b_norm not in _normalizar(u_nombre) and b_norm not in _normalizar(accion_legible):
                    continue

            fecha = _formato_fecha(a.registrado_en)
            detalle = a.nota or a.entidad or ""
            ws.append([
                fecha,
                u_nombre,
                accion_legible,
                tipo,
                detalle,
            ])

    # Estilos de encabezado y congelamiento de fila
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    _FORMULA_PREFIXES = ("=", "+", "-", "@")
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(_FORMULA_PREFIXES):
                    cell.data_type = "s"

    temp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    temp_path = temp.name
    temp.close()
    wb.save(temp_path)

    return temp_path
