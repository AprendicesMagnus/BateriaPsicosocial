"""Llena las plantillas Excel de informes de una evaluación y las entrega en un ZIP.

Fuentes de cada dato:
- Ficha de datos (F1..F21): respuestas del cuestionario FICHA_DATOS.
- Intralaboral (dimensiones, dominios y total): resultados guardados por tabulacion.py
  (ya usan factores y baremos oficiales de docs/referencia_intralaboral_A_B.json).
- Extralaboral (7 dimensiones + total) y estrés: se calculan aquí desde las respuestas
  con la calificación oficial del Manual del Ministerio de la Protección Social (2010),
  porque en la BD cada uno es una sola dimensión con baremos genéricos.
- Puntaje total general (intralaboral + extralaboral): factores y baremos de la referencia oficial.
"""
import re
import unicodedata
from copy import copy
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from sqlalchemy.orm import Session, joinedload

from app.core.config import get_settings
from app.core.encryption import descifrar_json
from app.db.intralaboral_oficial import cargar_referencia
from app.models.evaluation import Evaluacion, EvaluacionParticipante, Respuesta, ResultadoDimension
from app.models.organization import Organizacion
from app.models.survey import CuestionarioVersion, Dimension, Pregunta
from app.models.user import Usuario

CARPETA_PLANTILLAS = Path(__file__).resolve().parents[1] / "templates" / "informes_excel"

ZONA_COLOMBIA = timezone(timedelta(hours=-5))

NIVELES = ["SIN_RIESGO", "BAJO", "MEDIO", "ALTO", "MUY_ALTO"]
NIVEL_TEXTO = {
    "SIN_RIESGO": "Sin riesgo",
    "BAJO": "Riesgo bajo",
    "MEDIO": "Riesgo medio",
    "ALTO": "Riesgo alto",
    "MUY_ALTO": "Riesgo muy alto",
    "SIN_CLASIFICAR": "Sin clasificar",
}

# Etiquetas de respuesta según el valor que guarda el frontend (utils/encuestas.js)
ETIQUETAS_RESPUESTA = {
    "ESTRES": {4: "Siempre", 3: "Casi siempre", 2: "A veces", 1: "Nunca"},
    "EXTRALABORAL": {5: "Siempre", 4: "Casi siempre", 3: "Algunas veces", 2: "Casi nunca", 1: "Nunca"},
    "INTRALABORAL": {4: "Siempre", 3: "Casi siempre", 2: "Algunas veces", 1: "Casi nunca", 0: "Nunca"},
}

# ---------------------------------------------------------------------------
# Calificación oficial del cuestionario extralaboral (Manual, tablas 11 a 18)
# ---------------------------------------------------------------------------
# Ítems que puntúan Siempre=4 ... Nunca=0; el resto puntúa Siempre=0 ... Nunca=4.
EXTRA_ITEMS_DIRECTOS = {2, 3, 6, 24, 26, 28, 30, 31}
EXTRA_DIMENSIONES = {
    "TIEMPO": ([14, 15, 16, 17], 16),
    "FAMILIA": ([22, 25, 27], 12),
    "COMUNICACION": ([18, 19, 20, 21, 23], 20),
    "ECONOMICA": ([29, 30, 31], 12),
    "VIVIENDA": ([5, 6, 7, 8, 9, 10, 11, 12, 13], 36),
    "INFLUENCIA": ([24, 26, 28], 12),
    "DESPLAZAMIENTO": ([1, 2, 3, 4], 16),
}
EXTRA_FACTOR_TOTAL = 124
# Límite superior de cada nivel (Sin riesgo, Bajo, Medio, Alto); por encima es Muy alto.
EXTRA_BAREMOS = {
    "A": {
        "TIEMPO": (6.3, 25.0, 37.5, 50.0),
        "FAMILIA": (8.3, 25.0, 33.3, 50.0),
        "COMUNICACION": (0.9, 10.0, 20.0, 30.0),
        "ECONOMICA": (8.3, 25.0, 33.3, 50.0),
        "VIVIENDA": (5.6, 11.1, 13.9, 22.2),
        "INFLUENCIA": (8.3, 16.7, 25.0, 41.7),
        "DESPLAZAMIENTO": (0.9, 12.5, 25.0, 43.8),
        "TOTAL": (11.3, 16.9, 22.6, 29.0),
    },
    "B": {
        "TIEMPO": (6.3, 25.0, 37.5, 50.0),
        "FAMILIA": (8.3, 25.0, 33.3, 50.0),
        "COMUNICACION": (5.0, 15.0, 25.0, 35.0),
        "ECONOMICA": (16.7, 25.0, 41.7, 50.0),
        "VIVIENDA": (5.6, 11.1, 16.7, 27.8),
        "INFLUENCIA": (0.9, 16.7, 25.0, 41.7),
        "DESPLAZAMIENTO": (0.9, 12.5, 25.0, 43.8),
        "TOTAL": (12.9, 17.7, 24.2, 32.3),
    },
}

# ---------------------------------------------------------------------------
# Calificación oficial del cuestionario de estrés (3a versión, tablas 4 a 6)
# ---------------------------------------------------------------------------
# Puntos de cada opción (Siempre, Casi siempre, A veces, Nunca) por grupo de ítems
ESTRES_PUNTOS = {
    **{n: (9, 6, 3, 0) for n in (1, 2, 3, 9, 13, 14, 15, 23, 24)},
    **{n: (6, 4, 2, 0) for n in (4, 5, 6, 10, 11, 16, 17, 18, 19, 25, 26, 27, 28)},
    **{n: (3, 2, 1, 0) for n in (7, 8, 12, 20, 21, 22, 29, 30, 31)},
}
# (primer ítem, último ítem, peso) para el puntaje bruto ponderado
ESTRES_BLOQUES = [(1, 8, 4), (9, 12, 3), (13, 22, 2), (23, 31, 1)]
ESTRES_FACTOR = 61.16
ESTRES_BAREMOS = {
    "A": (7.8, 12.6, 17.7, 25.0),
    "B": (6.5, 11.8, 17.0, 23.4),
}

# ---------------------------------------------------------------------------
# Nombres que aparecen en las plantillas -> código interno del resultado
# ---------------------------------------------------------------------------
ALIAS_RESULTADOS = {
    # Dimensiones intralaborales
    "CARACTERISTICAS DEL LIDERAZGO": "INTRA:LIDERAZGO",
    "RELACIONES SOCIALES EN EL TRABAJO": "INTRA:RELACIONES_SOCIALES",
    "RETROALIMENTACION DEL DESEMPENO": "INTRA:RETROALIMENTACION",
    "RELACION CON LOS COLABORADORES": "INTRA:RELACION_COLABORADORES",
    "CLARIDAD DE ROL": "INTRA:CLARIDAD_ROL",
    "CAPACITACION": "INTRA:CAPACITACION",
    "PARTICIPACION Y MANEJO DEL CAMBIO": "INTRA:PARTICIPACION_CAMBIO",
    "OPORTUNIDADES PARA EL USO Y DESARROLLO DE HABILIDADES Y CONOCIMIENTOS": "INTRA:OPORTUNIDADES_HABILIDADES",
    "OPORTUNIDADES DE DESARROLLO Y USO DE HABILIDADES Y CONOCIMIENTOS": "INTRA:OPORTUNIDADES_HABILIDADES",
    "CONTROL Y AUTONOMIA SOBRE EL TRABAJO": "INTRA:CONTROL_AUTONOMIA",
    "DEMANDAS AMBIENTALES Y DE ESFUERZO FISICO": "INTRA:DEMANDAS_AMBIENTALES",
    "DEMANDAS EMOCIONALES": "INTRA:DEMANDAS_EMOCIONALES",
    "DEMANDAS CUANTITATIVAS": "INTRA:DEMANDAS_CUANTITATIVAS",
    "INFLUENCIA DEL TRABAJO SOBRE EL ENTORNO EXTRALABORAL": "INTRA:INFLUENCIA_EXTRALABORAL",
    "EXIGENCIAS DE RESPONSABILIDAD DEL CARGO": "INTRA:RESPONSABILIDAD_CARGO",
    "DEMANDAS DE CARGA MENTAL": "INTRA:CARGA_MENTAL",
    "CONSISTENCIA DEL ROL": "INTRA:CONSISTENCIA_ROL",
    "DEMANDAS DE LA JORNADA DE TRABAJO": "INTRA:JORNADA",
    "RECOMPENSAS DERIVADAS DE LA PERTENENCIA A LA ORGANIZACION Y DEL TRABAJO QUE SE REALIZA": "INTRA:RECOMPENSAS_PERTENENCIA",
    "RECONOCIMIENTO Y COMPENSACION": "INTRA:RECONOCIMIENTO_COMPENSACION",
    # Dominios intralaborales
    "LIDERAZGO Y RELACIONES SOCIALES EN EL TRABAJO": "INTRA:DOM_LIDERAZGO",
    "CONTROL SOBRE EL TRABAJO": "INTRA:DOM_CONTROL",
    "DEMANDAS DEL TRABAJO": "INTRA:DOM_DEMANDAS",
    "RECOMPENSAS": "INTRA:DOM_RECOMPENSAS",
    # Dimensiones extralaborales
    "TIEMPO FUERA DEL TRABAJO": "EXTRA:TIEMPO",
    "RELACIONES FAMILIARES": "EXTRA:FAMILIA",
    "COMUNICACION Y RELACIONES INTERPERSONALES": "EXTRA:COMUNICACION",
    "SITUACION ECONOMICA DEL GRUPO FAMILIAR": "EXTRA:ECONOMICA",
    "CARACTERISTICAS DE LA VIVIENDA Y DE SU ENTORNO": "EXTRA:VIVIENDA",
    "INFLUENCIA DEL ENTORNO EXTRALABORAL SOBRE EL TRABAJO": "EXTRA:INFLUENCIA",
    "DESPLAZAMIENTO VIVIENDA TRABAJO VIVIENDA": "EXTRA:DESPLAZAMIENTO",
    # Totales
    "INTRALABORAL": "INTRA:TOTAL_CUESTIONARIO",
    "EXTRALABORAL": "EXTRA:TOTAL",
    "ESTRES": "ESTRES:TOTAL",
    "INTRALABORAL EXTRALABORAL": "GENERAL:TOTAL",
}

# Encabezados de datos sociodemográficos -> campo del registro
COLUMNAS_SOCIO = {
    "FECHA DE APLICACION": "fecha",
    "DOCUMENTO": "documento",
    "NOMBRE": "nombre1",
    "NOMBRE 2": "nombre2",
    "APELLIDO": "apellido1",
    "APELLIDO 2": "apellido2",
    "GENERO": "F2",
    "ANO DE NACIMIENTO": "F3",
    "ESTADO CIVIL": "F4",
    "ESCOLARIDAD": "F5",
    "OCUPACION O PROFESION": "F6",
    "CIUDAD MUNICIPIO RESIDENCIA": "F7",
    "DEPARTAMENTO RESIDENCIA": "F8",
    "ESTRATO": "F9",
    "TIPO DE VIVIENDA": "F10",
    "PERSONAS A CARGO": "F11",
    "CIUDAD MUNICIPIO DONDE TRABAJA": "F12",
    "DEPARTAMENTO DONDE TRABAJA": "F13",
    "ANTIGUEDAD EN LA EMPRESA": "F14",
    "NOMBRE DEL CARGO": "F15",
    "TIPO DE CARGO": "F16",
    "ANTIGUEDAD EN EL CARGO": "F17",
    "DEPARTAMENTO O SECCION DE LA EMPRESA DONDE TRABAJA": "F18",
    "TIPO DE CONTRATO": "F19",
    "HORAS DIARIAS DE TRABAJO": "F20",
    "TIPO DE SALARIO": "F21",
}

# Bloques de la hoja "Sociodemografica" de Resultados -> (campo, se compara como número)
BLOQUES_SOCIO = {
    "GENERO": ("F2", False),
    "RANGO DE EDAD": ("edad", True),
    "ESTADO CIVIL": ("F4", False),
    "ESTUDIO": ("F5", False),
    "VIVIENDA": ("F10", False),
    "ESTRATO": ("F9", False),
    "PERSONAS A CARGO": ("F11", True),
    "CARGO": ("F16", False),
    "ANTIGUEDAD EN EL CARGO": ("F17", True),
    "ANTIGUEDAD EN LA EMPRESA": ("F14", True),
    "TIPO DE SALARIO": ("F21", False),
    "TIPO DE CONTRATO": ("F19", False),
    "HORAS TRABAJADAS": ("F20", True),
}
# Diferencias de redacción entre las opciones de la ficha (frontend) y las plantillas
EQUIVALENCIAS_TEXTO = {
    "NO SE": "NO SABE",
    "TECNICO TECNOLOGICO INCOMPLETO": "TECNICO TECNOLOGO INCOMPLETO",
    "TECNICO TECNOLOGICO COMPLETO": "TECNICO TECNOLOGO COMPLETO",
}

HOJAS_ESTADISTICAS_A = ["Total_Intralaboral_A", "Dominios_A", "Dimensiones_A", "Extralaboral_A", "Estres_A"]


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------
def _clave(valor) -> str:
    """Texto comparable: sin tildes, sin paréntesis, solo letras/números en mayúscula."""
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"\(.*?(\)|$)", " ", texto)
    texto = re.sub(r"[^0-9A-Za-z]+", " ", texto)
    return texto.strip().upper()


def _resolver(etiqueta) -> str | None:
    """Código de resultado para una etiqueta de la plantilla (o None si no es un resultado)."""
    clave = _clave(etiqueta)
    clave = re.sub(r"^(DOMINIO|DIMENSION)\s+", "", clave)
    clave = re.sub(r"\s+FORMA [AB]$", "", clave)
    if clave.endswith(" EXTRALABORAL") and clave not in ALIAS_RESULTADOS:
        clave = clave[: -len(" EXTRALABORAL")]
    return ALIAS_RESULTADOS.get(clave)


def _escribir(hoja, fila, columna, valor, formato=None):
    celda = hoja.cell(fila, columna)
    if isinstance(celda, MergedCell):
        return
    celda.value = valor
    if formato:
        celda.number_format = formato


def _entero(valor) -> int | None:
    try:
        return int(str(valor).strip())
    except (TypeError, ValueError):
        return None


def _numero_item(codigo: str) -> int | None:
    coincidencia = re.search(r"(\d+)$", codigo or "")
    return int(coincidencia.group(1)) if coincidencia else None


def _clasificar(puntaje, limites) -> str | None:
    """Nivel según límites superiores (Sin riesgo, Bajo, Medio, Alto)."""
    if puntaje is None:
        return None
    for nivel, maximo in zip(NIVELES, limites):
        if puntaje <= maximo:
            return nivel
    return "MUY_ALTO"


def _limites_baremos(baremos) -> tuple:
    """Convierte baremos [{nivel, maximo, orden}] a límites superiores ordenados."""
    ordenados = sorted(baremos, key=lambda b: b["orden"] if isinstance(b, dict) else b.orden)
    return tuple((b["maximo"] if isinstance(b, dict) else b.maximo) for b in ordenados)[:4]


def _fecha_local(valor):
    if valor is None:
        return None
    if valor.tzinfo is not None:
        valor = valor.astimezone(ZONA_COLOMBIA)
    return valor.date()


def _separar_nombre(texto: str) -> tuple[str, str, str, str]:
    partes = str(texto or "").split()
    if len(partes) >= 4:
        return partes[0], " ".join(partes[1:-2]), partes[-2], partes[-1]
    if len(partes) == 3:
        return partes[0], "", partes[1], partes[2]
    if len(partes) == 2:
        return partes[0], "", partes[1], ""
    return (partes[0] if partes else ""), "", "", ""


# ---------------------------------------------------------------------------
# Cálculos oficiales extralaboral y estrés
# ---------------------------------------------------------------------------
def _calcular_extralaboral(respuestas: dict, forma: str) -> tuple[dict, float | None]:
    """Devuelve ({codigo: (puntaje, nivel)}, puntaje_bruto_total)."""
    puntos = {}
    for numero, valor in respuestas.items():
        if valor is None or not 1 <= valor <= 5:
            continue
        # El frontend guarda Siempre=5 ... Nunca=1
        puntos[numero] = valor - 1 if numero in EXTRA_ITEMS_DIRECTOS else 5 - valor

    resultados = {}
    baremos = EXTRA_BAREMOS[forma]
    for codigo, (items, factor) in EXTRA_DIMENSIONES.items():
        # Si falta algún ítem de la dimensión, el manual la considera no válida.
        if any(n not in puntos for n in items):
            continue
        puntaje = round(100.0 * sum(puntos[n] for n in items) / factor, 1)
        resultados[f"EXTRA:{codigo}"] = (puntaje, _clasificar(puntaje, baremos[codigo]))

    bruto_total = None
    if all(n in puntos for n in range(1, 32)):
        bruto_total = float(sum(puntos.values()))
        puntaje = round(100.0 * bruto_total / EXTRA_FACTOR_TOTAL, 1)
        resultados["EXTRA:TOTAL"] = (puntaje, _clasificar(puntaje, baremos["TOTAL"]))
    return resultados, bruto_total


def _calcular_estres(respuestas: dict, forma: str) -> tuple[float, str] | None:
    if any(respuestas.get(n) not in (1, 2, 3, 4) for n in range(1, 32)):
        return None
    bruto = 0.0
    for desde, hasta, peso in ESTRES_BLOQUES:
        # El frontend guarda Siempre=4, Casi siempre=3, A veces=2, Nunca=1
        puntos = [ESTRES_PUNTOS[n][4 - respuestas[n]] for n in range(desde, hasta + 1)]
        bruto += (sum(puntos) / len(puntos)) * peso
    puntaje = round(100.0 * bruto / ESTRES_FACTOR, 1)
    return puntaje, _clasificar(puntaje, ESTRES_BAREMOS[forma])


# ---------------------------------------------------------------------------
# Carga de datos
# ---------------------------------------------------------------------------
def _forma_por_cargo(tipo_cargo) -> str | None:
    clave = _clave(tipo_cargo)
    if clave.startswith(("JEFATURA", "PROFESIONAL")):
        return "A"
    if clave.startswith(("AUXILIAR", "OPERARIO")):
        return "B"
    return None


def _participantes(db: Session, organizacion_id, evaluacion_id=None, area_id=None) -> list[EvaluacionParticipante]:
    """Participantes con respuestas de una evaluación o, sin ella, de todas las de la empresa.

    Cada enlace de pacientes crea su propia evaluación, por eso el informe de empresa las junta.
    Si un trabajador participó varias veces se toma su participación más reciente.
    """
    consulta = (
        db.query(EvaluacionParticipante)
        .join(Evaluacion, EvaluacionParticipante.evaluacion_id == Evaluacion.id)
        .options(joinedload(EvaluacionParticipante.trabajador), joinedload(EvaluacionParticipante.evaluacion))
        .filter(Evaluacion.organizacion_id == organizacion_id, EvaluacionParticipante.respuestas.any())
    )
    if evaluacion_id is not None:
        consulta = consulta.filter(Evaluacion.id == evaluacion_id)
    if area_id is not None:
        consulta = consulta.join(Usuario, Usuario.id == EvaluacionParticipante.trabajador_id).filter(
            Usuario.area_id == area_id
        )
    participantes = consulta.order_by(Evaluacion.creado_en.desc()).all()

    unicos = {}
    for participante in participantes:
        unicos.setdefault(participante.trabajador_id, participante)
    return list(unicos.values())


def _cargar_registros(db: Session, participantes: list[EvaluacionParticipante]) -> list[dict]:
    if not participantes:
        return []
    ids = [p.id for p in participantes]

    respuestas = (
        db.query(Respuesta.participante_id, Respuesta.valor_cifrado, Pregunta.codigo, CuestionarioVersion.codigo)
        .join(Pregunta, Respuesta.pregunta_id == Pregunta.id)
        .join(Dimension, Pregunta.dimension_id == Dimension.id)
        .join(CuestionarioVersion, Dimension.version_id == CuestionarioVersion.id)
        .filter(Respuesta.participante_id.in_(ids))
        .all()
    )
    resultados = (
        db.query(ResultadoDimension, Dimension.codigo, CuestionarioVersion.codigo)
        .join(Dimension, ResultadoDimension.dimension_id == Dimension.id)
        .join(CuestionarioVersion, Dimension.version_id == CuestionarioVersion.id)
        .filter(
            ResultadoDimension.participante_id.in_(ids),
            CuestionarioVersion.codigo.in_(["INTRALABORAL_A", "INTRALABORAL_B"]),
        )
        .all()
    )

    base = {
        p.id: {
            "participante": p,
            "ficha": {},
            "respuestas": {"ESTRES": {}, "EXTRALABORAL": {}, "INTRALABORAL": {}},
            "forma_intra": None,
            "puntajes": {},
            "intra_bruto_total": None,
        }
        for p in participantes
    }

    for participante_id, valor_cifrado, codigo_pregunta, codigo_version in respuestas:
        registro = base[participante_id]
        try:
            valor = descifrar_json(valor_cifrado)
        except (ValueError, TypeError):
            continue
        if codigo_version == "FICHA_DATOS":
            registro["ficha"][str(codigo_pregunta).strip().upper()] = "" if valor is None else str(valor).strip()
            continue
        numero = _numero_item(codigo_pregunta)
        entero = _entero(valor)
        if numero is None or entero is None:
            continue
        if codigo_version in ("INTRALABORAL_A", "INTRALABORAL_B"):
            registro["forma_intra"] = codigo_version[-1]
            # Respuestas antiguas en escala 1-5 (mismo ajuste que tabulacion.py)
            registro["respuestas"]["INTRALABORAL"][numero] = entero - 1 if entero > 4 else entero
        elif codigo_version in ("ESTRES", "EXTRALABORAL"):
            registro["respuestas"][codigo_version][numero] = entero

    for resultado, codigo_dimension, codigo_version in resultados:
        registro = base[resultado.participante_id]
        registro["forma_intra"] = codigo_version[-1]
        registro["puntajes"][f"INTRA:{codigo_dimension}"] = (resultado.puntaje_transformado, resultado.nivel)
        if codigo_dimension == "TOTAL_CUESTIONARIO":
            registro["intra_bruto_total"] = resultado.puntaje_bruto

    referencia = cargar_referencia()

    registros = []
    for registro in base.values():
        tiene_datos = registro["ficha"] or registro["puntajes"] or any(registro["respuestas"].values())
        if not tiene_datos:
            continue
        participante = registro["participante"]
        usuario = participante.trabajador
        ficha = registro["ficha"]

        forma = registro["forma_intra"] or _forma_por_cargo(ficha.get("F16")) or "A"
        registro["forma"] = forma
        registro["forma_aplicada"] = registro["forma_intra"]

        extra, extra_bruto = _calcular_extralaboral(registro["respuestas"]["EXTRALABORAL"], forma)
        registro["puntajes"].update(extra)
        estres = _calcular_estres(registro["respuestas"]["ESTRES"], forma)
        if estres:
            registro["puntajes"]["ESTRES:TOTAL"] = estres

        if registro["intra_bruto_total"] is not None and extra_bruto is not None:
            general = referencia[f"forma_{forma}"]["total_general_intra_extra"]
            puntaje = round(100.0 * (registro["intra_bruto_total"] + extra_bruto) / general["factor"], 1)
            registro["puntajes"]["GENERAL:TOTAL"] = (puntaje, _clasificar(puntaje, _limites_baremos(general["baremos"])))

        # Nombre: el de la cuenta; para pacientes por enlace (invitados) el de la ficha
        if usuario is not None and not getattr(usuario, "es_invitado", False) and usuario.nombre:
            nombres_cuenta = str(usuario.nombre).split()
            apellidos = str(usuario.apellido or "").split()
            nombres = (
                nombres_cuenta[0],
                " ".join(nombres_cuenta[1:]),
                apellidos[0] if apellidos else "",
                " ".join(apellidos[1:]),
            )
            nombre_completo = f"{usuario.nombre} {usuario.apellido or ''}".strip()
        else:
            nombre_completo = ficha.get("F1", "")
            nombres = _separar_nombre(nombre_completo)

        ano_nacimiento = _entero(ficha.get("F3"))
        fecha_referencia = (
            participante.fecha_fin
            or participante.fecha_inicio
            or participante.evaluacion.fecha_inicio
            or participante.evaluacion.creado_en
            or datetime.now(ZONA_COLOMBIA)
        )
        registro.update(
            {
                "nombre_completo": nombre_completo,
                "nombre1": nombres[0],
                "nombre2": nombres[1],
                "apellido1": nombres[2],
                "apellido2": nombres[3],
                "documento": (getattr(usuario, "numero_identificacion", None) or "") if usuario else "",
                "fecha": _fecha_local(participante.fecha_fin or participante.fecha_inicio),
                "edad": fecha_referencia.year - ano_nacimiento if ano_nacimiento else None,
            }
        )
        registros.append(registro)

    registros.sort(key=lambda r: _clave(r["nombre_completo"]))
    return registros


def _baremos_intralaborales(db: Session) -> dict:
    """{(forma, codigo_dimension): límites} de las versiones intralaborales vigentes."""
    dimensiones = (
        db.query(Dimension, CuestionarioVersion.codigo)
        .options(joinedload(Dimension.baremos))
        .join(CuestionarioVersion, Dimension.version_id == CuestionarioVersion.id)
        .filter(CuestionarioVersion.codigo.in_(["INTRALABORAL_A", "INTRALABORAL_B"]), CuestionarioVersion.vigente.is_(True))
        .all()
    )
    return {(version[-1], d.codigo): _limites_baremos(d.baremos) for d, version in dimensiones if d.baremos}


def _valor_ficha(registro: dict, campo: str):
    if campo in ("F14", "F17"):
        anos = _entero(registro["ficha"].get(campo))
        if anos is None:
            return registro["ficha"].get(campo, "")
        return "Menos de un año" if anos == 0 else anos
    if campo in ("F3", "F9", "F11", "F20"):
        numero = _entero(registro["ficha"].get(campo))
        return numero if numero is not None else registro["ficha"].get(campo, "")
    if campo.startswith("F"):
        return registro["ficha"].get(campo, "")
    return registro.get(campo, "")


# ---------------------------------------------------------------------------
# Estadísticas de grupo
# ---------------------------------------------------------------------------
class Estadisticas:
    def __init__(self, registros: list[dict], baremos_intra: dict):
        self.registros = registros
        self.baremos_intra = baremos_intra

    def del_grupo(self, forma: str) -> list[dict]:
        return [r for r in self.registros if r["forma"] == forma]

    def distribucion(self, registros: list[dict], codigo: str) -> tuple[list[int], int]:
        conteo = [0] * len(NIVELES)
        for registro in registros:
            nivel = registro["puntajes"].get(codigo, (None, None))[1]
            if nivel in NIVELES:
                conteo[NIVELES.index(nivel)] += 1
        return conteo, sum(conteo)

    def _limites(self, codigo: str, forma: str):
        grupo, nombre = codigo.split(":", 1)
        if grupo == "INTRA":
            return self.baremos_intra.get((forma, nombre))
        if grupo == "EXTRA":
            return EXTRA_BAREMOS[forma].get(nombre)
        if grupo == "ESTRES":
            return ESTRES_BAREMOS[forma]
        general = cargar_referencia()[f"forma_{forma}"]["total_general_intra_extra"]
        return _limites_baremos(general["baremos"])

    def promedio(self, forma: str, codigo: str) -> tuple[float | None, str | None]:
        """Puntaje promedio del grupo y su nivel según el baremo de esa forma."""
        puntajes = [
            r["puntajes"][codigo][0]
            for r in self.del_grupo(forma)
            if codigo in r["puntajes"] and r["puntajes"][codigo][0] is not None
        ]
        if not puntajes:
            return None, None
        promedio = round(sum(puntajes) / len(puntajes), 1)
        limites = self._limites(codigo, forma)
        return promedio, (_clasificar(promedio, limites) if limites else None)


def _texto_nivel(nivel) -> str:
    return NIVEL_TEXTO.get(nivel, nivel or "")


# ---------------------------------------------------------------------------
# Escritura de plantillas
# ---------------------------------------------------------------------------
def _limpiar_filas(hoja, fila_inicio: int) -> None:
    for fila in hoja.iter_rows(min_row=fila_inicio, max_row=hoja.max_row):
        for celda in fila:
            if not isinstance(celda, MergedCell):
                celda.value = None


def _copiar_estilo_fila(hoja, fila_origen: int, fila_destino: int) -> None:
    if fila_origen == fila_destino:
        return
    for columna in range(1, hoja.max_column + 1):
        origen = hoja.cell(fila_origen, columna)
        if origen.has_style:
            hoja.cell(fila_destino, columna)._style = copy(origen._style)


def _columnas_socio(hoja, fila_encabezado: int) -> dict[int, str]:
    columnas = {}
    for columna in range(1, hoja.max_column + 1):
        valor = hoja.cell(fila_encabezado, columna).value
        if valor is None:
            continue
        if str(valor).strip() == "#":
            columnas[columna] = "indice"
        elif _clave(valor) in COLUMNAS_SOCIO:
            columnas[columna] = COLUMNAS_SOCIO[_clave(valor)]
    return columnas


def _escribir_fila_socio(hoja, fila: int, indice: int, registro: dict, columnas: dict) -> None:
    for columna, campo in columnas.items():
        if campo == "indice":
            _escribir(hoja, fila, columna, indice)
        elif campo == "fecha":
            _escribir(hoja, fila, columna, registro["fecha"], "DD/MM/YYYY")
        else:
            _escribir(hoja, fila, columna, _valor_ficha(registro, campo))


def _escribir_sociodemografica(libro, registros: list[dict]) -> None:
    hoja = libro.worksheets[0]
    columnas = _columnas_socio(hoja, 1)
    _limpiar_filas(hoja, 2)
    for indice, registro in enumerate(registros, start=1):
        fila = indice + 1
        _copiar_estilo_fila(hoja, 2, fila)
        _escribir_fila_socio(hoja, fila, indice, registro, columnas)


def _escribir_hoja_resultados(hoja, registros: list[dict]) -> None:
    """Hoja 'Forma A' / 'Forma B': una fila por trabajador con ficha y todos sus puntajes."""
    columnas = _columnas_socio(hoja, 2)

    # Sección de cada columna según los títulos de la fila 1
    secciones = []
    for columna in range(1, hoja.max_column + 1):
        titulo = _clave(hoja.cell(1, columna).value)
        if titulo:
            secciones.append((columna, titulo))

    def seccion_de(columna):
        actual = ""
        for inicio, titulo in secciones:
            if columna >= inicio:
                actual = titulo
        return actual

    columnas_puntaje = {}
    for columna in range(1, hoja.max_column + 1):
        if columna in columnas:
            continue
        encabezado = hoja.cell(2, columna).value
        texto = _clave(encabezado)
        if not texto:
            continue
        tipo = "puntaje" if "PUNTAJE TRANSFORMADO" in _clave(str(encabezado).replace("(", " ")) else "nivel"
        seccion = seccion_de(columna)
        if "PUNTAJE TOTAL" in texto:
            if "EXTRALABORAL" in seccion:
                codigo = "EXTRA:TOTAL"
            elif "ESTRES" in seccion:
                codigo = "ESTRES:TOTAL"
            else:
                codigo = "INTRA:TOTAL_CUESTIONARIO"
        else:
            codigo = _resolver(encabezado)
            if codigo and "EXTRALABORAL" in seccion and not codigo.startswith("EXTRA:"):
                codigo = None
        if codigo:
            columnas_puntaje[columna] = (codigo, tipo)

    _limpiar_filas(hoja, 3)
    for indice, registro in enumerate(registros, start=1):
        fila = indice + 2
        _copiar_estilo_fila(hoja, 3, fila)
        _escribir_fila_socio(hoja, fila, indice, registro, columnas)
        for columna, (codigo, tipo) in columnas_puntaje.items():
            puntaje, nivel = registro["puntajes"].get(codigo, (None, None))
            if tipo == "puntaje":
                _escribir(hoja, fila, columna, puntaje, "0.0")
            else:
                _escribir(hoja, fila, columna, _texto_nivel(nivel) if nivel else None)


def _escribir_niveles(hoja, fila, columna_inicio, conteo, total, porcentaje):
    for desplazamiento, cantidad in enumerate(conteo):
        if porcentaje:
            _escribir(hoja, fila, columna_inicio + desplazamiento, (cantidad / total) if total else 0, "0.00%")
        else:
            _escribir(hoja, fila, columna_inicio + desplazamiento, cantidad)


def _escribir_hoja_niveles(hoja, registros: list[dict], estadisticas: Estadisticas) -> None:
    """Hojas Total_Intralaboral / Dominios / Dimensiones / Extralaboral / Estres por grupo."""
    modo_porcentaje = False
    # Solo se escriben filas debajo de un encabezado "Sin riesgo ..."; así los títulos
    # como "ESTRES FORMA A" no se confunden con filas de datos.
    tabla_izquierda = tabla_derecha = False
    for fila in range(1, hoja.max_row + 1):
        etiqueta_b = hoja.cell(fila, 2).value
        clave_b = _clave(etiqueta_b)
        if "PORCENTAJE" in clave_b:
            modo_porcentaje = True
        elif clave_b.startswith("POR CANTIDAD"):
            modo_porcentaje = False

        if clave_b == "TOTAL TRABAJADORES":
            _escribir(hoja, fila, 3, len(registros))
            continue
        if _clave(hoja.cell(fila, 3).value) == "SIN RIESGO":
            tabla_izquierda = True
        if _clave(hoja.cell(fila, 10).value) == "SIN RIESGO":
            tabla_derecha = True

        codigo = _resolver(etiqueta_b) if etiqueta_b and tabla_izquierda else None
        if codigo:
            conteo, total = estadisticas.distribucion(registros, codigo)
            _escribir_niveles(hoja, fila, 3, conteo, total, modo_porcentaje)

        # Tabla de porcentajes a la derecha (columna I con etiqueta, J..N con valores)
        etiqueta_i = hoja.cell(fila, 9).value
        codigo_i = _resolver(etiqueta_i) if etiqueta_i and tabla_derecha else None
        if codigo_i:
            conteo, total = estadisticas.distribucion(registros, codigo_i)
            _escribir_niveles(hoja, fila, 10, conteo, total, True)

    _escribir_frecuencias_estres(hoja, registros)


def _escribir_frecuencias_estres(hoja, registros: list[dict]) -> None:
    """Distribución Siempre / Casi siempre / A veces / Nunca por ítem del cuestionario de estrés."""
    fila_titulos = None
    for fila in range(1, min(hoja.max_row, 10) + 1):
        if _clave(hoja.cell(fila, 11).value) == "SIEMPRE":
            fila_titulos = fila
            break
    if fila_titulos is None:
        return
    numero = 0
    for fila in range(fila_titulos + 1, hoja.max_row + 1):
        if not hoja.cell(fila, 10).value:
            continue
        numero += 1
        valores = [r["respuestas"]["ESTRES"].get(numero) for r in registros]
        for desplazamiento, valor in enumerate((4, 3, 2, 1)):
            _escribir(hoja, fila, 11 + desplazamiento, sum(1 for v in valores if v == valor))


def _coincide_numero(etiqueta: str, numero: int) -> bool:
    clave = _clave(etiqueta)
    numeros = [int(n) for n in re.findall(r"\d+", clave)]
    if "MENOR DE EDAD" in clave:
        return numero < 18
    if clave.startswith("NINGUN"):
        return numero == 0
    if clave.startswith("MENOS DE"):
        return numero < (numeros[0] if numeros else 1)
    if clave.startswith("MAS DE") and numeros:
        return numero > numeros[0]
    if "Y MAS" in clave and numeros:
        return numero >= numeros[0]
    if len(numeros) >= 2:
        return numeros[0] <= numero <= numeros[1]
    if len(numeros) == 1:
        return numero == numeros[0]
    return False


def _coincide_texto(etiqueta: str, valor) -> bool:
    clave_valor = _clave(valor)
    if not clave_valor:
        return False
    clave_valor = EQUIVALENCIAS_TEXTO.get(clave_valor, clave_valor)
    clave_etiqueta = _clave(etiqueta)
    if clave_valor == clave_etiqueta:
        return True
    numero = _entero(valor)
    # Estrato: "1" frente a "Estrato 1"
    return numero is not None and re.findall(r"\d+", clave_etiqueta) == [str(numero)]


def _escribir_conteos_socio(hoja, registros: list[dict]) -> None:
    """Hoja 'Sociodemografica' de Resultados: cantidad y porcentaje por categoría."""
    fila = 1
    while fila <= hoja.max_row:
        if _clave(hoja.cell(fila, 2).value) != "CANTIDAD":
            fila += 1
            continue
        configuracion = BLOQUES_SOCIO.get(_clave(hoja.cell(fila, 1).value))
        categorias = []
        fila += 1
        while fila <= hoja.max_row and hoja.cell(fila, 1).value:
            categorias.append(fila)
            fila += 1
        if not configuracion:
            continue
        campo, numerico = configuracion
        valores = [r.get("edad") if campo == "edad" else r["ficha"].get(campo) for r in registros]
        valores = [v for v in valores if v not in (None, "")]
        for fila_categoria in categorias:
            etiqueta = hoja.cell(fila_categoria, 1).value
            if numerico:
                cantidad = sum(1 for v in valores if _entero(v) is not None and _coincide_numero(etiqueta, _entero(v)))
            else:
                cantidad = sum(1 for v in valores if _coincide_texto(etiqueta, v))
            _escribir(hoja, fila_categoria, 2, cantidad)
            _escribir(hoja, fila_categoria, 5, (cantidad / len(valores)) if valores else 0, "0.00%")


def _duplicar_para_forma_b(libro, nombre_hoja: str):
    original = libro[nombre_hoja]
    copia = libro.copy_worksheet(original)
    copia.title = "Forma B" if nombre_hoja == "Forma A" else re.sub(r"_A$", "_B", nombre_hoja)
    for fila in copia.iter_rows():
        for celda in fila:
            if isinstance(celda.value, str) and not isinstance(celda, MergedCell):
                celda.value = celda.value.replace("FORMA A", "FORMA B").replace("Forma A", "Forma B")
    return copia


def _escribir_resultados(libro, estadisticas: Estadisticas) -> None:
    grupo_a = estadisticas.del_grupo("A")
    grupo_b = estadisticas.del_grupo("B")

    _escribir_conteos_socio(libro["Sociodemografica"], estadisticas.registros)

    hojas_b = []
    if grupo_b:
        hojas_b = [_duplicar_para_forma_b(libro, nombre) for nombre in ["Forma A", *HOJAS_ESTADISTICAS_A]]

    _escribir_hoja_resultados(libro["Forma A"], grupo_a)
    for nombre in HOJAS_ESTADISTICAS_A:
        _escribir_hoja_niveles(libro[nombre], grupo_a, estadisticas)

    if hojas_b:
        _escribir_hoja_resultados(hojas_b[0], grupo_b)
        for hoja in hojas_b[1:]:
            _escribir_hoja_niveles(hoja, grupo_b, estadisticas)


def _escribir_consolidado(libro, registros: list[dict]) -> None:
    hoja = libro.worksheets[0]
    encabezados = {
        2: "DOCUMENTO", 3: "NOMBRE COMPLETO", 4: "NOMBRE DEL CARGO", 5: "FORMA APLICADA",
        6: "PUNTAJE", 7: "NIVEL DE RIESGO", 8: "PUNTAJE", 9: "NIVEL DE RIESGO", 10: "PUNTAJE", 11: "NIVEL DE RIESGO",
    }
    for columna, texto in encabezados.items():
        celda = hoja.cell(2, columna)
        if celda.value is None:
            celda._style = copy(hoja.cell(2, 1)._style)
            celda.value = texto

    columnas = [(6, "INTRA:TOTAL_CUESTIONARIO"), (8, "EXTRA:TOTAL"), (10, "ESTRES:TOTAL")]
    _limpiar_filas(hoja, 3)
    for indice, registro in enumerate(registros, start=1):
        fila = indice + 2
        _copiar_estilo_fila(hoja, 3, fila)
        forma = f"Forma {registro['forma_aplicada']}" if registro["forma_aplicada"] else ""
        _escribir(hoja, fila, 1, indice)
        _escribir(hoja, fila, 2, registro["documento"])
        _escribir(hoja, fila, 3, registro["nombre_completo"])
        _escribir(hoja, fila, 4, registro["ficha"].get("F15", ""))
        _escribir(hoja, fila, 5, forma)
        for columna, codigo in columnas:
            puntaje, nivel = registro["puntajes"].get(codigo, (None, None))
            _escribir(hoja, fila, columna, puntaje, "0.0")
            _escribir(hoja, fila, columna + 1, _texto_nivel(nivel) if nivel else None)


def _escribir_resumen(libro, estadisticas: Estadisticas, organizacion) -> None:
    hoja = libro.worksheets[0]
    _escribir(hoja, 1, 2, getattr(organizacion, "nombre", "") or "Empresa")

    for columna, forma in ((2, "A"), (12, "B")):
        texto = hoja.cell(4, columna).value
        if isinstance(texto, str):
            cantidad = len(estadisticas.del_grupo(forma))
            _escribir(hoja, 4, columna, re.sub(r"para los \d+ trabajadores", f"para los {cantidad} trabajadores", texto))

    def resumen(forma, codigo):
        promedio, nivel = estadisticas.promedio(forma, codigo)
        if promedio is None:
            return None
        return f"{_texto_nivel(nivel)} ({promedio:.1f})" if nivel else f"{promedio:.1f}"

    seccion = "GENERAL"
    for fila in range(1, hoja.max_row + 1):
        clave_b = _clave(hoja.cell(fila, 2).value)
        if clave_b.startswith("CONDICIONES INTRALABORALES"):
            seccion = "INTRA"
        elif clave_b.startswith("CONDICIONES EXTRALABORALES"):
            seccion = "EXTRA"
        elif clave_b.startswith("CONDICIONES DE ESTRES"):
            seccion = "ESTRES"
        elif clave_b.startswith("NIVEL DE RIESGO"):
            break

        if seccion == "GENERAL":
            codigo = _resolver(hoja.cell(fila, 2).value)
            if not codigo:
                continue
            # (puntaje A, nivel A, puntaje B, nivel B)
            columnas = (7, 9, 17, 19) if codigo == "GENERAL:TOTAL" else (5, 8, 15, 18)
            for forma, (col_puntaje, col_nivel) in (("A", columnas[:2]), ("B", columnas[2:])):
                promedio, nivel = estadisticas.promedio(forma, codigo)
                _escribir(hoja, fila, col_puntaje, promedio, "0.0")
                _escribir(hoja, fila, col_nivel, _texto_nivel(nivel) if nivel else None)
        elif seccion == "INTRA":
            # Dominio en D (valores en G / I) y dimensión en K (valores en Q / S)
            for col_etiqueta, col_a, col_b in ((4, 7, 9), (11, 17, 19)):
                codigo = _resolver(hoja.cell(fila, col_etiqueta).value)
                if codigo and codigo.startswith("INTRA:"):
                    _escribir(hoja, fila, col_a, resumen("A", codigo))
                    _escribir(hoja, fila, col_b, resumen("B", codigo))
        elif seccion in ("EXTRA", "ESTRES"):
            col_etiqueta = 4 if seccion == "EXTRA" else 2
            codigo = _resolver(hoja.cell(fila, col_etiqueta).value)
            if codigo and codigo.split(":")[0] == seccion:
                _escribir(hoja, fila, 9, resumen("A", codigo))
                _escribir(hoja, fila, 11, resumen("B", codigo))


def _escribir_grupo(libro, registros: list[dict]) -> None:
    """Grupo_A / Grupo_B: una columna por persona con la respuesta de cada pregunta."""
    hoja = libro.worksheets[0]
    columna_inicio = 5  # Columna E = "Persona 1"
    ultima_plantilla = hoja.max_column

    for desplazamiento, registro in enumerate(registros):
        columna = columna_inicio + desplazamiento
        if columna > ultima_plantilla:
            # Más personas que columnas en la plantilla: se agrega una columna con el mismo formato
            for fila in range(1, hoja.max_row + 1):
                origen = hoja.cell(fila, ultima_plantilla)
                if origen.has_style and not isinstance(hoja.cell(fila, columna), MergedCell):
                    hoja.cell(fila, columna)._style = copy(origen._style)
            letra_origen = hoja.cell(1, ultima_plantilla).column_letter
            letra_nueva = hoja.cell(1, columna).column_letter
            hoja.column_dimensions[letra_nueva].width = hoja.column_dimensions[letra_origen].width
            _escribir(hoja, 4, columna, f"Persona {desplazamiento + 1}")

        for fila in range(5, hoja.max_row + 1):
            seccion = _clave(hoja.cell(fila, 1).value)
            numero = _entero(hoja.cell(fila, 2).value)
            pregunta = _clave(hoja.cell(fila, 3).value)
            if seccion == "DATOS":
                if pregunta == "NOMBRE":
                    _escribir(hoja, fila, columna, registro["nombre_completo"])
                elif pregunta.startswith("ID"):
                    _escribir(hoja, fila, columna, registro["documento"])
                continue
            if numero is None or seccion not in ETIQUETAS_RESPUESTA:
                continue
            valor = registro["respuestas"][seccion].get(numero)
            _escribir(hoja, fila, columna, ETIQUETAS_RESPUESTA[seccion].get(valor))


# ---------------------------------------------------------------------------
# Punto de entrada
# ---------------------------------------------------------------------------
def _cargar_plantilla(nombre: str):
    ruta = CARPETA_PLANTILLAS / nombre
    if not ruta.is_file():
        raise FileNotFoundError(f"No se encontró la plantilla: {nombre}")
    return load_workbook(ruta)


def _a_bytes(libro) -> bytes:
    memoria = BytesIO()
    libro.save(memoria)
    libro.close()
    return memoria.getvalue()


def generar_zip_informes_empresa(
    db: Session, organizacion_id, evaluacion_id=None, area_id=None, aplicar_minimo_anonimato=False
) -> BytesIO:
    """ZIP con las plantillas llenas; sin evaluacion_id incluye todas las evaluaciones de la empresa."""
    organizacion = db.query(Organizacion).filter(Organizacion.id == organizacion_id).first()
    if organizacion is None:
        raise ValueError("La empresa seleccionada no existe.")

    if evaluacion_id is not None:
        existe = (
            db.query(Evaluacion.id)
            .filter(Evaluacion.id == evaluacion_id, Evaluacion.organizacion_id == organizacion_id)
            .first()
        )
        if existe is None:
            raise ValueError("La evaluación no existe o no pertenece a la empresa seleccionada.")

    registros = _cargar_registros(db, _participantes(db, organizacion_id, evaluacion_id, area_id))

    if aplicar_minimo_anonimato:
        minimo = get_settings().min_grupo_anonimato
        if len(registros) < minimo:
            raise ValueError(
                f"No se puede generar el informe agrupado. El grupo tiene {len(registros)} participante(s) completados, "
                f"lo cual es inferior al mínimo requerido de anonimato ({minimo} persona(s))."
            )

    if not registros:
        raise ValueError("Todavía no hay cuestionarios respondidos para generar los informes.")

    estadisticas = Estadisticas(registros, _baremos_intralaborales(db))
    archivos = {}

    libro = _cargar_plantilla("Sociodemografica.xlsx")
    _escribir_sociodemografica(libro, registros)
    archivos["Sociodemografica.xlsx"] = _a_bytes(libro)

    libro = _cargar_plantilla("Resultados.xlsx")
    _escribir_resultados(libro, estadisticas)
    archivos["Resultados.xlsx"] = _a_bytes(libro)

    libro = _cargar_plantilla("Consolidado trabajadores.xlsx")
    _escribir_consolidado(libro, registros)
    archivos["Consolidado trabajadores.xlsx"] = _a_bytes(libro)

    libro = _cargar_plantilla("Resumen estadístico.xlsx")
    _escribir_resumen(libro, estadisticas, organizacion)
    archivos["Resumen estadístico.xlsx"] = _a_bytes(libro)

    for forma in ("A", "B"):
        grupo = estadisticas.del_grupo(forma)
        if grupo:
            libro = _cargar_plantilla(f"Grupo_{forma}.xlsx")
            _escribir_grupo(libro, grupo)
            archivos[f"Respuestas Grupo {forma}.xlsx"] = _a_bytes(libro)

    salida = BytesIO()
    with ZipFile(salida, "w", ZIP_DEFLATED) as zip_file:
        for nombre, contenido in archivos.items():
            zip_file.writestr(nombre, contenido)
    salida.seek(0)
    return salida
