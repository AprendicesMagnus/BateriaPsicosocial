import io
import uuid
from datetime import datetime, timezone
import openpyxl

from app.models.audit import Auditoria
from app.models.user import Usuario


def test_exportar_auditoria_etiquetas_precisas(client, db_session):
    res_admin = client.post(
        "/api/auth/login",
        json={"email": "admin@magnussig.com", "password": "Admin1234"},
    )
    assert res_admin.status_code == 200
    headers = {"Authorization": f"Bearer {res_admin.json()['token']}"}

    admin_user = db_session.query(Usuario).filter(Usuario.email == "admin@magnussig.com").first()
    assert admin_user is not None

    ahora = datetime.now(timezone.utc)
    db_session.add_all([
        Auditoria(
            id=uuid.uuid4(),
            usuario_id=admin_user.id,
            accion="INFORME_INDIVIDUAL_GENERADO",
            nota="etiqueta-ind-3f",
            registrado_en=ahora,
        ),
        Auditoria(
            id=uuid.uuid4(),
            usuario_id=admin_user.id,
            accion="INFORME_AGRUPADO_GENERADO",
            nota="etiqueta-agr-3f",
            registrado_en=ahora,
        ),
    ])
    db_session.commit()

    res = client.get("/api/panel/exportar/auditoria", headers=headers)
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb.active

    filas = list(ws.iter_rows(min_row=2, values_only=True))
    fila_ind = next((f for f in filas if f[4] == "etiqueta-ind-3f"), None)
    fila_agr = next((f for f in filas if f[4] == "etiqueta-agr-3f"), None)

    assert fila_ind is not None, "No se encontró fila con nota 'etiqueta-ind-3f'"
    assert fila_agr is not None, "No se encontró fila con nota 'etiqueta-agr-3f'"

    assert fila_ind[2] == "generó un informe individual"
    assert fila_agr[2] == "generó un informe agrupado"


def test_exportar_auditoria_buscar_etiqueta_agrupado(client, db_session):
    res_admin = client.post(
        "/api/auth/login",
        json={"email": "admin@magnussig.com", "password": "Admin1234"},
    )
    assert res_admin.status_code == 200
    headers = {"Authorization": f"Bearer {res_admin.json()['token']}"}

    admin_user = db_session.query(Usuario).filter(Usuario.email == "admin@magnussig.com").first()
    assert admin_user is not None

    ahora = datetime.now(timezone.utc)
    nota_ind = f"etiqueta-ind-buscar-{uuid.uuid4().hex[:6]}"
    nota_agr = f"etiqueta-agr-buscar-{uuid.uuid4().hex[:6]}"

    db_session.add_all([
        Auditoria(
            id=uuid.uuid4(),
            usuario_id=admin_user.id,
            accion="INFORME_INDIVIDUAL_GENERADO",
            nota=nota_ind,
            registrado_en=ahora,
        ),
        Auditoria(
            id=uuid.uuid4(),
            usuario_id=admin_user.id,
            accion="INFORME_AGRUPADO_GENERADO",
            nota=nota_agr,
            registrado_en=ahora,
        ),
    ])
    db_session.commit()

    # Buscar con texto legible
    res = client.get("/api/panel/exportar/auditoria?buscar=generó+un+informe+agrupado", headers=headers)
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb.active

    filas = list(ws.iter_rows(min_row=2, values_only=True))
    notas = [f[4] for f in filas]

    assert nota_agr in notas, "La fila agrupada debe aparecer en la búsqueda"
    assert nota_ind not in notas, "La fila individual no debe aparecer en la búsqueda de agrupado"
