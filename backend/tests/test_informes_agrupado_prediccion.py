import os
import uuid
import openpyxl
import pytest

from app.core.nit_utils import calcular_digito_verificador_nit
from app.db.seed import sembrar_todos_los_cuestionarios
from app.models.evaluation import Evaluacion, EvaluacionParticipante, ResultadoDimension
from app.models.organization import Area, Organizacion
from app.models.survey import CuestionarioVersion, Dimension
from app.models.user import Usuario
from app.services.informes import _perfiles_visibles


@pytest.fixture(autouse=True)
def setup_seed(db_session):
    sembrar_todos_los_cuestionarios(db_session)
    db_session.commit()


@pytest.fixture
def datos_prueba_agrupado(client, db_session, crear_paciente):
    # 1. Login admin
    res_admin = client.post(
        "/api/auth/login",
        json={"email": "admin@magnussig.com", "password": "Admin1234"},
    )
    assert res_admin.status_code == 200, res_admin.text
    headers_admin = {"Authorization": f"Bearer {res_admin.json()['token']}"}

    # 2. Crear Organización y Áreas
    nit_base = f"9{str(uuid.uuid4().int)[:8]}"
    nit_rnd = f"{nit_base}-{calcular_digito_verificador_nit(nit_base)}"
    org = Organizacion(id=uuid.uuid4(), nombre="Empresa Prediccion Agrupada SAS", nit=nit_rnd)
    db_session.add(org)
    db_session.commit()

    area_ops = Area(id=uuid.uuid4(), organizacion_id=org.id, nombre="Operaciones")
    area_ventas = Area(id=uuid.uuid4(), organizacion_id=org.id, nombre="Ventas")
    db_session.add_all([area_ops, area_ventas])
    db_session.commit()

    # 3. Obtener versión con dimensiones sembradas
    dim = db_session.query(Dimension).first()
    assert dim is not None, "Debe existir al menos una dimensión sembrada"
    cuestionario = (
        db_session.query(CuestionarioVersion)
        .filter(CuestionarioVersion.id == dim.version_id)
        .first()
    )
    assert cuestionario is not None

    dimensiones = (
        db_session.query(Dimension)
        .filter(Dimension.version_id == cuestionario.id)
        .all()
    )

    admin_user = (
        db_session.query(Usuario)
        .filter(Usuario.email == "admin@magnussig.com")
        .first()
    )
    assert admin_user is not None

    evaluacion = Evaluacion(
        id=uuid.uuid4(),
        organizacion_id=org.id,
        evaluador_id=admin_user.id,
        version_id=cuestionario.id,
        nombre="Evaluación Predictiva Agrupada Test 2026",
    )
    db_session.add(evaluacion)
    db_session.commit()

    # 4. Crear 5 participantes para Operaciones con puntaje 45.0 (Riesgo Moderado)
    participantes_ops_ids = [crear_paciente(org.id, area_id=area_ops.id)[0] for _ in range(5)]

    def _crear_participante_con_resultados(trabajador_id):
        part = EvaluacionParticipante(
            id=uuid.uuid4(),
            evaluacion_id=evaluacion.id,
            trabajador_id=trabajador_id,
            estado="COMPLETADA",
        )
        db_session.add(part)
        db_session.flush()

        for d in dimensiones:
            res_dim = ResultadoDimension(
                participante_id=part.id,
                dimension_id=d.id,
                puntaje_bruto=45.0,
                puntaje_transformado=45.0,
                nivel="MEDIO",
            )
            db_session.add(res_dim)

    for tid in participantes_ops_ids:
        _crear_participante_con_resultados(tid)

    db_session.commit()

    archivos_a_limpiar = []

    yield {
        "headers_admin": headers_admin,
        "evaluacion_id": str(evaluacion.id),
        "area_ops_id": str(area_ops.id),
        "archivos_a_limpiar": archivos_a_limpiar,
    }

    for ruta in archivos_a_limpiar:
        if os.path.exists(ruta):
            try:
                os.remove(ruta)
            except OSError:
                pass


def test_informe_agrupado_excel_con_analisis_predictivo(client, datos_prueba_agrupado):
    headers = datos_prueba_agrupado["headers_admin"]
    eval_id = datos_prueba_agrupado["evaluacion_id"]
    area_id = datos_prueba_agrupado["area_ops_id"]

    res = client.post(
        "/api/informes/agrupado",
        headers=headers,
        json={
            "evaluacionId": eval_id,
            "areaId": area_id,
            "formato": "EXCEL",
        },
    )
    assert res.status_code == 200, res.text
    data = res.json()

    ruta_archivo = data.get("rutaArchivo")
    assert ruta_archivo and os.path.exists(ruta_archivo)
    datos_prueba_agrupado["archivos_a_limpiar"].append(ruta_archivo)

    wb = openpyxl.load_workbook(ruta_archivo)
    assert "Análisis predictivo" in wb.sheetnames

    ws = wb["Análisis predictivo"]
    encontrado = any("Perfil de Riesgo Moderado" in str(cell.value) for row in ws.rows for cell in row)
    assert encontrado, "No se encontró 'Perfil de Riesgo Moderado' en la hoja de Análisis predictivo"


def test_informe_agrupado_pdf_con_analisis_predictivo_y_verificacion(client, datos_prueba_agrupado):
    headers = datos_prueba_agrupado["headers_admin"]
    eval_id = datos_prueba_agrupado["evaluacion_id"]
    area_id = datos_prueba_agrupado["area_ops_id"]

    res = client.post(
        "/api/informes/agrupado",
        headers=headers,
        json={
            "evaluacionId": eval_id,
            "areaId": area_id,
            "formato": "PDF",
        },
    )
    assert res.status_code == 200, res.text
    data = res.json()

    codigo_verif = data.get("codigoVerificacion")
    assert codigo_verif and len(codigo_verif.strip()) > 0

    hash_sha256 = data.get("hashSha256")
    assert hash_sha256 and len(hash_sha256) == 64

    ruta_archivo = data.get("rutaArchivo")
    assert ruta_archivo and os.path.exists(ruta_archivo)
    datos_prueba_agrupado["archivos_a_limpiar"].append(ruta_archivo)

    with open(ruta_archivo, "rb") as f:
        cabecera = f.read(4)
        assert cabecera == b"%PDF"

    # Verificar endpoint público
    res_verif = client.get(f"/api/informes/verificar/{codigo_verif}")
    assert res_verif.status_code == 200, res_verif.text
    data_verif = res_verif.json()
    assert data_verif["valido"] is True
    assert data_verif["tipo"] == "AGRUPADO"


def test_informe_agrupado_resiliencia_falla_predictiva(client, datos_prueba_agrupado, monkeypatch):
    headers = datos_prueba_agrupado["headers_admin"]
    eval_id = datos_prueba_agrupado["evaluacion_id"]
    area_id = datos_prueba_agrupado["area_ops_id"]

    def _falla_simulada(*args, **kwargs):
        raise RuntimeError("Error en clusterización K-Means")

    monkeypatch.setattr(
        "app.services.informes.prediccion_service.analizar_riesgo_predictivo",
        _falla_simulada,
    )

    res = client.post(
        "/api/informes/agrupado",
        headers=headers,
        json={
            "evaluacionId": eval_id,
            "areaId": area_id,
            "formato": "EXCEL",
        },
    )
    assert res.status_code == 200, res.text
    data = res.json()

    ruta_archivo = data.get("rutaArchivo")
    assert ruta_archivo and os.path.exists(ruta_archivo)
    datos_prueba_agrupado["archivos_a_limpiar"].append(ruta_archivo)

    wb = openpyxl.load_workbook(ruta_archivo)
    assert "Análisis predictivo" not in wb.sheetnames
    assert "Resumen Agrupado" in wb.sheetnames


def test_perfiles_visibles_regla_anonimato():
    prediccion = {
        "perfilesClusterKMeans": [
            {"etiqueta": "Cluster A", "numTrabajadores": 5},
            {"etiqueta": "Cluster B", "numTrabajadores": 2},
        ]
    }
    visibles, omitidos = _perfiles_visibles(prediccion, minimo=5)
    assert len(visibles) == 1
    assert visibles[0]["numTrabajadores"] == 5
    assert omitidos == 1
