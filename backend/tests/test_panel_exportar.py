import io
import uuid
from datetime import datetime, timezone
import openpyxl

from app.core.nit_utils import calcular_digito_verificador_nit
from app.models.audit import Auditoria
from app.models.commercial import Compra
from app.models.organization import Organizacion
from app.models.user import Usuario


def test_exportar_tablas_encabezados(client):
    # 1. Login Admin
    res_admin = client.post(
        "/api/auth/login",
        json={"email": "admin@magnussig.com", "password": "Admin1234"},
    )
    assert res_admin.status_code == 200, res_admin.text
    headers = {"Authorization": f"Bearer {res_admin.json()['token']}"}

    encabezados_esperados = {
        "usuarios": [
            "Nombre",
            "Correo",
            "Rol",
            "Empresa",
            "Registro",
            "Estado",
            "Último acceso",
            "Número de resolución",
            "Resolución verificada",
        ],
        "empresas": [
            "Empresa",
            "NIT",
            "Responsable SST",
            "Encuestas completadas",
            "Evaluaciones usadas",
            "Evaluaciones compradas",
        ],
        "compras": [
            "Fecha",
            "Empresa",
            "Cantidad",
            "Medio de pago",
            "Valor",
            "Estado",
        ],
        "auditoria": [
            "Fecha",
            "Usuario",
            "Acción",
            "Tipo",
            "Detalle",
        ],
    }

    for tabla, cabeceras in encabezados_esperados.items():
        res = client.get(f"/api/panel/exportar/{tabla}", headers=headers)
        assert res.status_code == 200, f"Fallo al exportar {tabla}: {res.text}"
        assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        primera_fila = [cell.value for cell in ws[1]]
        assert primera_fila == cabeceras, f"Encabezados incorrectos en {tabla}: {primera_fila} != {cabeceras}"


def test_exportar_compras_mas_de_50(client, db_session):
    res_admin = client.post(
        "/api/auth/login",
        json={"email": "admin@magnussig.com", "password": "Admin1234"},
    )
    assert res_admin.status_code == 200
    headers = {"Authorization": f"Bearer {res_admin.json()['token']}"}

    admin_user = db_session.query(Usuario).filter(Usuario.email == "admin@magnussig.com").first()
    assert admin_user is not None

    nit_base = f"9{str(uuid.uuid4().int)[:8]}"
    nit_rnd = f"{nit_base}-{calcular_digito_verificador_nit(nit_base)}"
    org = Organizacion(id=uuid.uuid4(), nombre="Empresa Compras Masivas SAS", nit=nit_rnd)
    db_session.add(org)
    db_session.commit()

    # Crear 55 compras
    compras = []
    for _ in range(55):
        compras.append(
            Compra(
                id=uuid.uuid4(),
                usuario_id=admin_user.id,
                organizacion_id=org.id,
                cantidad=1,
                subtotal=10000.0,
                iva=1900.0,
                total=11900.0,
                estado="PAGADA",
                creado_en=datetime.now(timezone.utc),
            )
        )
    db_session.add_all(compras)
    db_session.commit()

    res = client.get("/api/panel/exportar/compras", headers=headers)
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb.active

    filas = list(ws.rows)
    # Al menos 1 encabezado + 55 compras recién creadas
    assert len(filas) > 51, f"Se esperaban más de 50 compras en el archivo Excel, pero hay {len(filas) - 1}"


def test_exportar_sin_permisos(client, db_session, crear_paciente):
    # Sin token
    res_sin_token = client.get("/api/panel/exportar/usuarios")
    assert res_sin_token.status_code in (401, 403)

    # Con rol no superadmin
    nit_base = f"9{str(uuid.uuid4().int)[:8]}"
    nit_rnd = f"{nit_base}-{calcular_digito_verificador_nit(nit_base)}"
    org = Organizacion(id=uuid.uuid4(), nombre="Empresa Paciente SAS", nit=nit_rnd)
    db_session.add(org)
    db_session.commit()

    _pid, email_p = crear_paciente(org.id, password="Paciente123!")
    res_login = client.post(
        "/api/auth/login",
        json={"email": email_p, "password": "Paciente123!"},
    )
    assert res_login.status_code == 200
    headers_paciente = {"Authorization": f"Bearer {res_login.json()['token']}"}

    res_no_admin = client.get("/api/panel/exportar/usuarios", headers=headers_paciente)
    assert res_no_admin.status_code == 403


def test_exportar_tabla_inexistente(client):
    res_admin = client.post(
        "/api/auth/login",
        json={"email": "admin@magnussig.com", "password": "Admin1234"},
    )
    assert res_admin.status_code == 200
    headers = {"Authorization": f"Bearer {res_admin.json()['token']}"}

    res = client.get("/api/panel/exportar/inexistente", headers=headers)
    assert res.status_code == 400
    assert "no válida" in res.text


def test_exportar_auditoria_filtro_accesos(client, db_session):
    res_admin = client.post(
        "/api/auth/login",
        json={"email": "admin@magnussig.com", "password": "Admin1234"},
    )
    assert res_admin.status_code == 200
    headers = {"Authorization": f"Bearer {res_admin.json()['token']}"}

    admin_user = db_session.query(Usuario).filter(Usuario.email == "admin@magnussig.com").first()

    # Agregar explícitamente un registro de tipo LOGIN (Accesos) y otro de tipo CREAR_USUARIO (Usuarios)
    db_session.add(Auditoria(
        id=uuid.uuid4(),
        usuario_id=admin_user.id,
        accion="LOGIN",
        registrado_en=datetime.now(timezone.utc),
    ))
    db_session.add(Auditoria(
        id=uuid.uuid4(),
        usuario_id=admin_user.id,
        accion="CREAR_USUARIO",
        registrado_en=datetime.now(timezone.utc),
    ))
    db_session.commit()

    res = client.get("/api/panel/exportar/auditoria?filtro=Accesos", headers=headers)
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb.active

    filas = list(ws.rows)
    assert len(filas) >= 2, "Debe haber al menos una fila de accesos"

    # La columna 4 (índice 3) corresponde a 'Tipo'
    for fila in filas[1:]:
        tipo_celda = fila[3].value
        assert tipo_celda == "Accesos", f"Se esperaba tipo 'Accesos' pero se obtuvo '{tipo_celda}'"


def test_exportar_formula_injection_protegida(client, db_session):
    res_admin = client.post(
        "/api/auth/login",
        json={"email": "admin@magnussig.com", "password": "Admin1234"},
    )
    assert res_admin.status_code == 200
    headers = {"Authorization": f"Bearer {res_admin.json()['token']}"}

    nit_base = f"9{str(uuid.uuid4().int)[:8]}"
    nit_rnd = f"{nit_base}-{calcular_digito_verificador_nit(nit_base)}"
    org = Organizacion(id=uuid.uuid4(), nombre="=1+1", nit=nit_rnd)
    db_session.add(org)
    db_session.commit()

    res = client.get("/api/panel/exportar/empresas", headers=headers)
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.value == "=1+1":
                assert cell.data_type != "f", "La celda con '=1+1' no debe interpretarse como fórmula"
                return
    assert False, "No se encontró la celda con '=1+1'"


def test_exportar_auditoria_buscar_texto_legible(client, db_session):
    res_admin = client.post(
        "/api/auth/login",
        json={"email": "admin@magnussig.com", "password": "Admin1234"},
    )
    assert res_admin.status_code == 200
    headers = {"Authorization": f"Bearer {res_admin.json()['token']}"}

    admin_user = db_session.query(Usuario).filter(Usuario.email == "admin@magnussig.com").first()
    db_session.add(Auditoria(
        id=uuid.uuid4(),
        usuario_id=admin_user.id,
        accion="LOGIN",
        registrado_en=datetime.now(timezone.utc),
    ))
    db_session.commit()

    res = client.get("/api/panel/exportar/auditoria?buscar=inici%C3%B3+sesi%C3%B3n", headers=headers)
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb.active
    filas = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(filas) >= 1
    for fila in filas:
        assert fila[2] == "inició sesión", f"Columna Acción debe ser 'inició sesión', fue '{fila[2]}'"


def test_exportar_registra_auditoria_exportar_panel(client, db_session):
    from app.models.audit import Auditoria as Aud
    res_admin = client.post(
        "/api/auth/login",
        json={"email": "admin@magnussig.com", "password": "Admin1234"},
    )
    assert res_admin.status_code == 200
    headers = {"Authorization": f"Bearer {res_admin.json()['token']}"}

    antes = db_session.query(Aud).filter(Aud.accion == "EXPORTAR_PANEL", Aud.nota == "usuarios").count()
    client.get("/api/panel/exportar/usuarios", headers=headers)
    db_session.expire_all()
    despues = db_session.query(Aud).filter(Aud.accion == "EXPORTAR_PANEL", Aud.nota == "usuarios").count()
    assert despues > antes
